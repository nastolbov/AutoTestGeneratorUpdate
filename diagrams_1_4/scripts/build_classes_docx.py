# -*- coding: utf-8 -*-
"""Сборка docx-справочника классов программы из /tmp/api.json (см. extract_api.py).

Для каждого пакета -> каждого класса (включая вложенные и перечисления) выводит
назначение, наследование, таблицу полей [Доступ | Тип | Имя | Описание] и таблицу
методов [Доступ | Возвращает | Метод(параметры) | Описание]. Включаются все члены,
в т. ч. private. Оформление — как в дипломе (Times New Roman, поля ГОСТ, Table Grid).
"""
import os, re, json
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..")
API = "/tmp/api.json"
OUT = os.path.join(ROOT, "Описание_классов.docx")
FONT = "Times New Roman"
USABLE_CM = 16.5
PKG_ORDER = ["model", "parser", "generator", "data", "common", "ui"]

# ----------------------------------------------------------------------------- описания
PKG_RU = {
    "model": "Объектная модель метаданных предметной области, получаемая при разборе XML-файла.",
    "parser": "Разбор XML-файла метаданных и построение объектной модели предметной области.",
    "generator": "Генерация тестового проекта, объектов страниц, тест-классов и отчётов, запуск прогона.",
    "data": "Хранение результатов прогонов и доступ к ним в базе отчётов.",
    "common": "Вспомогательные компоненты общего назначения.",
    "ui": "Графический интерфейс пользователя (JavaFX).",
}

CLASS_RU = {
    "model.AppModel": "Корневая модель метаданных: агрегирует сущности, поиски и сведения о подсистеме, разобранные из XML.",
    "model.Association": "Ассоциация между сущностями модели: роли, кратности и связанный элемент.",
    "model.AttrType": "Тип значения свойства: строка, число, дата, дата-время.",
    "model.EntityClassifier": "Классификатор сущностей: относит каждую сущность к основной, дочерней или справочнику.",
    "model.EntityClassifier.Classification": "Результат классификации одной сущности (вид, причина, родительская сущность).",
    "model.EntityKind": "Вид сущности по результату классификации: основная, дочерняя, справочник.",
    "model.EntityObject": "Сущность предметной области: группы свойств, операции, ассоциации и признаки поведения.",
    "model.Modifier": "Модификатор операции — дополнительный признак её поведения.",
    "model.ModifyType": "Тип изменения записи (создание, редактирование, удаление и т. п.).",
    "model.Operation": "Операция сущности (действие бизнес-логики) с параметрами и модификаторами.",
    "model.OperationParam": "Параметр операции сущности.",
    "model.Property": "Свойство (атрибут) сущности: тип значения, обязательность, маска и прочие характеристики.",
    "model.PropertyGroup": "Группа свойств сущности — логический раздел формы со списком свойств.",
    "model.Search": "Описание поиска (фильтра) сущности с набором параметров и выводимых результатов.",
    "model.SearchParam": "Параметр поиска — критерий фильтрации.",
    "model.SearchResult": "Результат поиска: набор отображаемых свойств.",
    "model.SearchResultProperty": "Свойство, выводимое в результатах поиска.",
    "model.TestCaseResult": "Результат выполнения одного теста: статус, длительность, шаги и сведения об ошибке.",
    "model.TestCaseResult.StepTiming": "Тайминг отдельного шага теста: имя шага и его длительность.",
    "model.TestRunResult": "Результат прогона набора тестов: сводные показатели и перечень результатов тестов.",
    "parser.EntityParser": "Разбирает XML-описание сущности в объект EntityObject.",
    "parser.PropertyGroupParser": "Разбирает группы свойств и свойства сущности из XML.",
    "parser.SearchParser": "Разбирает описания поисков сущности из XML.",
    "parser.StaxUtils": "Вспомогательные методы потокового разбора XML (StAX).",
    "parser.XmlModelParser": "Точка входа разбора: строит модель AppModel из XML-файла метаданных.",
    "parser.XmlNamespaces": "Константы пространств имён XML, используемых в модели метаданных.",
    "generator.PageObjectWriter": "Генерирует класс объекта страницы (Page Object) для сущности.",
    "generator.RunReportWriter": "Формирует отчёты о прогоне тестов в форматах HTML и CSV.",
    "generator.TestClassWriter": "Генерирует тестовый класс с набором проверок для сущности.",
    "generator.TestConfig": "Параметры генерации и прогона: пути, адрес сайта, учётные данные, тип сайта и др.",
    "generator.TestDataFactory": "Формирует тестовые данные (значения свойств) по их типам и маскам.",
    "generator.TestGenerator": "Создаёт структуру тест-проекта и генерирует Page Object'ы и тест-классы по модели.",
    "generator.TestRunner": "Запускает сгенерированные автотесты (Maven) и собирает результаты прогона.",
    "data.DatabaseConnection": "Управляет соединением с базой данных отчётов.",
    "data.ReportDao": "Доступ к данным отчётов: чтение и запись сведений о прогонах.",
    "data.SchemaInitializer": "Создаёт схему базы данных отчётов (таблицы) при первом запуске.",
    "data.TestCaseDao": "Доступ к данным результатов отдельных тестов.",
    "data.TestRunDao": "Доступ к данным прогонов тестов.",
    "data.TestRunDao.RunRow": "Строка списка прогонов — запись о прогоне для отображения.",
    "common.JavaFileWriter": "Буфер генерации Java-кода с управлением отступами и записью в файл.",
    "common.ParserException": "Исключение, возникающее при разборе XML-модели метаданных.",
    "common.Transliterator": "Транслитерация русских наименований в латиницу для идентификаторов кода.",
    "ui.App": "JavaFX-приложение: инициализация и запуск главного окна.",
    "ui.Launcher": "Класс запуска приложения (точка входа, обход ограничений модульности JavaFX).",
    "ui.MainController": "Контроллер главного окна: связывает элементы интерфейса с разбором, генерацией и прогоном.",
    "ui.MainController.EntityNode": "Узел дерева сущностей в интерфейсе.",
    "ui.MainController.TestCaseRow": "Строка таблицы результатов тестов в интерфейсе.",
}

FIELD_RU = {
    "name": "системное имя", "guid": "глобальный идентификатор (GUID)",
    "title": "отображаемое наименование (заголовок)", "caption": "подпись (заголовок)",
    "orderNumber": "порядковый номер", "flagDisplay": "признак отображения",
    "valueType": "тип значения", "type": "тип", "kind": "вид",
    "basePackage": "базовый Java-пакет генерируемого кода", "connection": "соединение с базой данных",
    "featureName": "имя подсистемы (функциональности)", "searchGuid": "идентификатор поиска (GUID)",
    "params": "список параметров", "stereoType": "стереотип", "dmodule": "модуль (dmodule)",
    "required": "признак обязательности", "mask": "маска ввода", "properties": "список свойств",
    "propertyGroups": "список групп свойств", "className": "имя класса", "methodName": "имя метода",
    "passed": "число успешных тестов", "skipped": "число пропущенных тестов",
    "failed": "число проваленных тестов", "total": "общее число тестов",
    "durationMs": "длительность, мс", "baseUrl": "базовый адрес тестируемого сайта",
    "testLevel": "уровень тестирования", "entities": "список сущностей", "searches": "список поисков",
    "operations": "список операций", "modifiers": "список модификаторов",
    "roleA": "роль A ассоциации", "roleB": "роль B ассоциации",
    "roleACaption": "подпись роли A", "roleBCaption": "подпись роли B",
    "associationId": "идентификатор ассоциации", "reason": "причина классификации",
    "parentEntity": "родительская сущность", "parentGrid": "родительская таблица",
    "status": "статус", "error": "сообщение об ошибке", "steps": "шаги теста",
    "value": "значение", "login": "логин", "password": "пароль", "siteType": "тип сайта",
    "xmlFile": "файл метаданных (XML)", "outputDir": "каталог генерации",
    "stepName": "имя шага", "feature": "подсистема (функциональность)",
}

# точечные переопределения описаний методов: "пакет.Класс#метод" -> текст
METHOD_RU = {
    "parser.XmlModelParser#parse": "Разбирает XML-файл метаданных и возвращает модель AppModel.",
    "generator.TestGenerator#generate": "Генерирует тестовый проект по модели: инфраструктуру, Page Object'ы и тест-классы.",
    "generator.TestRunner#run": "Запускает сгенерированные автотесты и возвращает результат прогона.",
    "common.Transliterator#transliterate": "Транслитерирует русский текст в латиницу.",
}

ENUM_CONST_RU = {
    "AttrType.STRING": "Строковое значение", "AttrType.DECIMAL": "Числовое значение",
    "AttrType.DATE": "Дата", "AttrType.DATETIME": "Дата и время",
    "EntityKind.PRIMARY": "Основная сущность", "EntityKind.CHILD": "Дочерняя сущность",
    "EntityKind.REFERENCE_DICTIONARY": "Справочник",
    "ModifyType.INSERT": "Создание записи", "ModifyType.UPDATE": "Изменение записи",
    "ModifyType.DELETE": "Удаление записи", "ModifyType.LOGICAL_EDIT": "Логическое редактирование",
    "ModifyType.ARCHIVE": "Архивирование (перенос в архив)",
}

# суффиксы имён элементов интерфейса -> русское название элемента
UI_SUFFIX = [("ComboBox", "список"), ("CheckBox", "флажок"), ("TextField", "поле ввода"),
             ("PasswordField", "поле пароля"), ("TextArea", "область текста"), ("Button", "кнопка"),
             ("TableView", "таблица"), ("TableColumn", "столбец таблицы"), ("TreeView", "дерево"),
             ("Field", "поле"), ("Btn", "кнопка"), ("Table", "таблица"), ("Column", "столбец"),
             ("Tree", "дерево"), ("Label", "надпись"), ("Combo", "список"), ("Check", "флажок"),
             ("Area", "область"), ("Pane", "панель"), ("Box", "блок")]

VERB = {
    "write": "Генерирует код:", "parse": "Разбирает", "generate": "Генерирует",
    "build": "Формирует", "create": "Создаёт", "make": "Создаёт", "read": "Читает",
    "load": "Загружает", "save": "Сохраняет", "store": "Сохраняет", "insert": "Сохраняет (вставляет) в БД",
    "select": "Выбирает из БД", "update": "Обновляет", "delete": "Удаляет", "remove": "Удаляет",
    "add": "Добавляет", "find": "Находит", "resolve": "Определяет", "compute": "Вычисляет",
    "calc": "Вычисляет", "classify": "Классифицирует", "extract": "Извлекает", "collect": "Собирает",
    "map": "Сопоставляет", "apply": "Применяет", "handle": "Обрабатывает", "process": "Обрабатывает",
    "on": "Обработчик события:", "init": "Инициализирует", "initialize": "Инициализирует",
    "run": "Выполняет", "execute": "Выполняет", "open": "Открывает", "close": "Закрывает",
    "show": "Отображает", "refresh": "Обновляет", "select": "Выбирает", "to": "Преобразует в",
    "from": "Создаёт из", "format": "Форматирует", "escape": "Экранирует", "esc": "Экранирует",
    "csv": "Формирует CSV:", "indent": "Увеличивает отступ", "unindent": "Уменьшает отступ",
    "transliterate": "Транслитерирует", "snake": "Преобразует в snake_case", "looks": "Проверяет, что",
    "main": "Точка входа", "now": "Возвращает текущие дату и время", "today": "Возвращает текущую дату",
    "folder": "Формирует имя папки", "short": "Сокращает", "clean": "Очищает", "link": "Связывает",
    "ensure": "Гарантирует наличие", "skip": "Пропускает", "start": "Запускает", "stop": "Останавливает",
    "name": "Формирует имя", "attr": "Возвращает атрибут", "grid": "Работа с таблицей:",
}
ACCESS_WORDS = {"public", "protected", "private", "пакетный"}


def split_access(a):
    toks = a.split()
    acc = toks[0] if toks and toks[0] in ACCESS_WORDS else "пакетный"
    mods = " ".join(t for t in toks if t not in ACCESS_WORDS)
    return acc, mods


def humanize(name):
    return re.sub(r"(?<=[a-z0-9])(?=[A-Z])", " ", name).strip()


def field_ru(fname):
    if fname in FIELD_RU:
        return FIELD_RU[fname]
    for suf, ru in UI_SUFFIX:
        if fname.endswith(suf) and len(fname) > len(suf):
            base = humanize(fname[:-len(suf)]).lower()
            return f"{ru} «{base}»" if base else ru
    return humanize(fname).lower()


def accessor_field(n):
    for p in ("get", "set"):
        if n.startswith(p) and len(n) > len(p) and n[len(p)].isupper():
            return n[len(p)].lower() + n[len(p) + 1:]
    for p in ("is", "has"):
        if n.startswith(p) and len(n) > len(p) and n[len(p)].isupper():
            return n[len(p)].lower() + n[len(p) + 1:]
    return None


def gloss(name):
    m = re.match(r"[a-z]+", name)
    v = m.group(0) if m else ""
    rest = humanize(name[len(v):]).strip()
    base = VERB.get(v)
    if base is None:
        return None
    return (base + " " + rest).strip() if rest else base


def describe_field(f):
    if f["doc"]:
        return f["doc"]
    d = field_ru(f["name"])
    return d[:1].upper() + d[1:]


def describe_method(qual_pkg, m):
    key = qual_pkg + "#" + m["name"]
    if key in METHOD_RU:
        return METHOD_RU[key]
    if m["doc"]:
        return m["doc"]
    if m["ctor"]:
        return "Конструктор с параметрами." if m["params"] else "Конструктор по умолчанию."
    fn = accessor_field(m["name"])
    if fn is not None:
        d = field_ru(fn); d = d[:1].lower() + d[1:]
        if m["name"].startswith("set"):
            return "Устанавливает " + d + "."
        if m["name"].startswith(("is", "has")):
            return "Возвращает признак «" + d + "»."
        return "Возвращает " + d + "."
    g = gloss(m["name"])
    return g if g else humanize(m["name"]).capitalize()


# ----------------------------------------------------------------------------- оформление
doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT; st.font.size = Pt(14)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.line_spacing = 1.15; st.paragraph_format.space_after = Pt(0)
sec = doc.sections[0]
sec.left_margin = Cm(3); sec.right_margin = Cm(1.5); sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)


def _f(run, size=14, bold=False, italic=False):
    run.font.name = FONT; run.font.size = Pt(size); run.bold = bold; run.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def title(text, size, before=10, after=8, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph(); p.alignment = align
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    _f(p.add_run(text), size, bold=True); return p


def body(text, size=14, italic=False, after=4):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.25); p.paragraph_format.space_after = Pt(after)
    _f(p.add_run(text), size, italic=italic); return p


def _shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexc)
    tcPr.append(sh)


def _ct(cell, text, bold=False, size=11, align=None, mono=False):
    cell.text = ""; p = cell.paragraphs[0]
    if align:
        p.alignment = align
    p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    if mono:
        run.font.name = "Consolas"; run.font.size = Pt(size); run.bold = bold
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    else:
        _f(run, size, bold=bold)


def _fixed_layout(t):
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)


def grid_table(headers, rows, widths, mono_cols=()):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    _fixed_layout(t)
    for j, h in enumerate(headers):
        _ct(t.rows[0].cells[j], h, bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(t.rows[0].cells[j], "D9D9D9")
    # повтор шапки на каждой странице
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true"); trPr.append(th)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            _ct(t.rows[i].cells[j], val, size=11, mono=(j in mono_cols))
    for i in range(len(rows) + 1):
        for j, wd in enumerate(widths):
            t.rows[i].cells[j].width = Cm(wd)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def sig(m):
    ps = ", ".join((p["type"] + " " + p["name"]) for p in m["params"])
    return m["name"] + "(" + ps + ")"


KIND_RU = {"class": "класс", "interface": "интерфейс", "enum": "перечисление"}

# ----------------------------------------------------------------------------- содержимое
api = json.load(open(API, encoding="utf-8"))

title("Описание классов программного обеспечения", 16, before=0, after=10,
      align=WD_ALIGN_PARAGRAPH.CENTER)
body("В разделе приведено описание всех классов программы, сгруппированных по пакетам. Для каждого "
     "класса указаны его вид, модификатор доступа и наследование, а также таблицы полей и методов. "
     "В таблицах для каждого члена приведены модификатор доступа, тип (для метода — тип возвращаемого "
     "значения), имя с параметрами и краткое описание. Приведены все члены, включая закрытые (private). "
     "Описание сформировано автоматически по исходному коду.")

cls_count = fld_count = mth_count = 0
for pkg in PKG_ORDER:
    types = api.get(pkg, [])
    title(f"Пакет ru.autotestgen.{pkg}", 14, before=14)
    body(PKG_RU.get(pkg, ""), italic=True)
    for r in types:
        cls_count += 1
        full = f"ru.autotestgen.{pkg}.{r['qualified']}"
        title(full, 13, before=10, after=2)
        # шапка класса
        head = KIND_RU[r["kind"]] + ", " + r["access"]
        if r["extends"]:
            head += f"; наследует {r['extends']}"
        if r["implements"]:
            head += f"; реализует {', '.join(r['implements'])}"
        if r["nested"]:
            head += f"; вложенные: {', '.join(r['nested'])}"
        body(head, size=12, italic=True, after=2)
        purpose = CLASS_RU.get(f"{pkg}.{r['qualified']}") or r["doc"]
        if purpose:
            body("Назначение. " + purpose, size=12, after=4)
        # перечисление: константы
        if r["kind"] == "enum" and r["enum_constants"]:
            rows = [[c["name"], ENUM_CONST_RU.get(f"{r['name']}.{c['name']}", c["doc"] or "")]
                    for c in r["enum_constants"]]
            grid_table(["Константа", "Описание"], rows, [4.0, 12.5], mono_cols=(0,))
        # поля
        if r["fields"]:
            rows = []
            for f in r["fields"]:
                acc, mods = split_access(f["access"])
                typ = (mods + " " if mods else "") + f["type"]
                rows.append([acc, typ, f["name"], describe_field(f)])
                fld_count += 1
            grid_table(["Доступ", "Тип", "Поле", "Описание"], rows,
                       [2.6, 4.4, 3.3, 6.2], mono_cols=(1, 2))
        # методы (конструкторы первыми)
        if r["methods"]:
            ms = sorted(r["methods"], key=lambda m: (0 if m["ctor"] else 1))
            rows = []
            for m in ms:
                acc, mods = split_access(m["access"])
                ret = "—" if m["ctor"] else ((mods + " " if mods else "") + m["returns"])
                rows.append([acc, ret, sig(m), describe_method(f"{pkg}.{r['qualified']}", m)])
                mth_count += 1
            grid_table(["Доступ", "Возвращает", "Метод (параметры)", "Описание"], rows,
                       [2.4, 3.0, 5.6, 5.5], mono_cols=(1, 2))

doc.save(OUT)
print("Сохранено:", OUT)
print(f"Классов(типов): {cls_count} | полей: {fld_count} | методов/конструкторов: {mth_count}")
