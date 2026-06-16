# -*- coding: utf-8 -*-
"""Сборка отдельного документа «Модульная структура программного средства и структурная
карта Константайна»: модуль = файл .java. Документ содержит модульную структуру (простые
блоки-файлы по пакетам) и построенную из неё структурную карту Константайна (связи модулей
по данным и управлению с куплетами). Стиль/нумерация — как в разделах 1.5 (Times New Roman,
поля ГОСТ, подпись таблицы в две строки, рисунок по центру). Нумерация рисунков/таблицы с 1.
Хелперы оформления скопированы из build_section_153_docx.py."""
import os, sys, subprocess
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import const_model as CM
import module_model as MM
import module_files_model as MF

DIA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
ROOT = os.path.join(DIA, "..")
OUT = os.path.join(ROOT, "Модульная_структура_и_карта_Константайна.docx")
FONT = "Times New Roman"
USABLE_W = 16.5
MAX_H = 23.0
TYPE_RU = {"module": "модуль", "library": "библиотека", "pojo": "класс данных",
           "enum": "перечисление", "exception": "исключение"}

# назначение модулей: активные модули/библиотеки — из module_model; классы данных — здесь
DESC = {m[0]: m[6] for m in MM.MODULES}
DESC.update({
    "AppModel": "Корневая модель метаданных: категория, список сущностей и поисков",
    "EntityObject": "Сущность предметной области: имя, группы свойств, ассоциации",
    "Association": "Ассоциация сущностей (роли, ссылки, признаки отображения)",
    "PropertyGroup": "Группа свойств сущности (форма или грид) и её операция",
    "Property": "Свойство-поле формы: имя, тип, маска, обязательность",
    "Operation": "CRUD-операция группы свойств и её модификаторы",
    "OperationParam": "Параметр метода CRUD-операции",
    "Modifier": "Модификатор операции (вид изменения данных)",
    "Search": "Поисковый блок (фильтр): параметры и описание результата",
    "SearchParam": "Параметр поиска: имя, тип значения, маска",
    "SearchResult": "Описание грида результатов поиска",
    "SearchResultProperty": "Колонка грида результатов поиска",
    "TestRunResult": "Результат прогона автотестов: итоги и список кейсов",
    "TestCaseResult": "Результат отдельного теста: статус, сообщение, скриншоты",
    "EntityKind": "Вид сущности: PRIMARY / CHILD / REFERENCE_DICTIONARY",
    "AttrType": "Тип атрибута: STRING / DECIMAL / DATE / DATETIME",
    "ModifyType": "Код изменения: INSERT / UPDATE / DELETE / LOGICAL_EDIT / ARCHIVE",
    "ParserException": "Исключение, сигнализирующее об ошибке разбора XML-метамодели",
})

doc = Document()
st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(14)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)
sec = doc.sections[0]
sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)            # A4
sec.left_margin = Cm(3); sec.right_margin = Cm(1.5); sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)


def _f(run, size=14, bold=False, italic=False):
    run.font.name = FONT; run.font.size = Pt(size); run.bold = bold; run.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def plural(n, one, few, many):
    n10, n100 = n % 10, n % 100
    if n10 == 1 and n100 != 11:
        return one
    if 2 <= n10 <= 4 and not 12 <= n100 <= 14:
        return few
    return many


def heading(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6); p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.5
    _f(p.add_run(text), 14, bold=False); return p


def body(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.75); p.paragraph_format.line_spacing = 1.5
    _f(p.add_run(text), 14); return p


def gap(n):
    for _ in range(n):
        p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.5
        _f(p.add_run(""), 14)


def figure(fname, num, caption):
    w, h = Image.open(os.path.join(DIA, fname)).size
    disp = USABLE_W
    if (h / w) * USABLE_W > MAX_H:
        disp = MAX_H * w / h
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5; p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(os.path.join(DIA, fname), width=Cm(disp))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.line_spacing = 1.5
    _f(c.add_run(f"Рисунок {num} – {caption}"), 14)
    gap(1)


def _shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexc); tcPr.append(sh)


def _ct(cell, text, bold=False):
    cell.text = ""; p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
    _f(p.add_run(text), 12, bold=bold)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true"); trPr.append(el)


def _no_split(row):
    trPr = row._tr.get_or_add_trPr(); el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true"); trPr.append(el)


def table(num, caption, headers, widths, rows):
    p1 = doc.add_paragraph(); p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.paragraph_format.line_spacing = 1.0; p1.paragraph_format.space_before = Pt(6)
    _f(p1.add_run(f"Таблица {num}"), 14)
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER; p2.paragraph_format.line_spacing = 1.0
    _f(p2.add_run(caption), 14)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.LEFT; t.autofit = False
    el = OxmlElement("w:tblLayout"); el.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(el)
    grid = t._tbl.find(qn("w:tblGrid"))
    for gc, wd in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(round(wd * 567))))
    for j, h in enumerate(headers):
        _ct(t.rows[0].cells[j], h, bold=True); _shade(t.rows[0].cells[j], "D9D9D9")
    _repeat_header(t.rows[0])
    for i, cols in enumerate(rows, 1):
        _no_split(t.rows[i])
        for j, val in enumerate(cols):
            _ct(t.rows[i].cells[j], val)
    for i in range(len(rows) + 1):
        for j, wd in enumerate(widths):
            t.rows[i].cells[j].width = Cm(wd)
    gap(2)


# ============================================================ содержимое
n_mod = len(CM.MODULES)
n_pkg = len({m[0] for m in CM.MODULES})
from collections import Counter
kc = Counter(MF.KIND[cls] for _p, _f2, cls, _r, _c in CM.MODULES)

heading("Модульная структура программного средства и структурная карта Константайна")
body("Информационная система AutoTestGenerator реализована на языке Java и организована по "
     "модульному принципу. В качестве модуля рассматривается отдельный файл исходного кода "
     f"(.java); всего система содержит {n_mod} {plural(n_mod, 'модуль', 'модуля', 'модулей')}, "
     f"распределённых по {n_pkg} пакетам: ui (интерфейс пользователя), parser (разбор XML-"
     "метамодели), model (модель данных предметной области), generator (генерация автотестов), "
     "data (хранилище отчётов) и common (служебные утилиты).")
body(f"По назначению модули делятся на {kc['module']} "
     f"{plural(kc['module'], 'функциональный модуль', 'функциональных модуля', 'функциональных модулей')} "
     f"(имеют собственную логику управления), {kc['library']} повторно используемых библиотек, "
     f"{kc['pojo']} классов "
     f"данных (контейнеры метаданных предметной области и результатов прогона), {kc['enum']} "
     f"перечисления (используются как управляющие признаки) и {kc['exception']} класс исключения.")

body("Модульная структура системы — разбиение программного средства на пакеты и модули (.java) — "
     "приведена на рисунке 1. Корневой узел соответствует программному средству в целом, "
     "пунктирные области — пакетам, прямоугольники — отдельным модулям (файлам исходного кода). "
     "Структура отражает функциональную декомпозицию: каждый пакет объединяет модули одного "
     "назначения, а сам модуль решает одну законченную подзадачу.")
figure("module_files_struct.png", 1, "Модульная структура программного средства")

body("На основе модульной структуры построена структурная карта Константайна (рисунок 2), "
     "которая показывает обращения между модулями и потоки данных в нотации Л. Константайна. "
     "Прямоугольником обозначен функциональный модуль, прямоугольником с двойной рамкой — "
     "повторно используемая библиотека, овалом — класс данных (POJO), пунктирным овалом — "
     "перечисление, пунктирным прямоугольником — класс исключения. Сплошная стрелка соответствует "
     "обращению одного модуля к другому; над стрелкой указаны куплеты связи: «↓○» — данные, "
     "передаваемые в вызываемый модуль, «↑○» — данные, возвращаемые вызывающему модулю, «↑●» — "
     "управляющий признак.")
body("Классы данных пакета model показаны на карте отдельными узлами. Связь «модуль → класс "
     "данных» с куплетами «↓○» означает, что модуль формирует (заполняет поля) этого класса: так, "
     "парсеры EntityParser, PropertyGroupParser и SearchParser создают объекты EntityObject, "
     "Property, Operation, Search и другие из разбираемого XML. Связь с куплетами «↑○» означает "
     "чтение полей класса данных модулем-потребителем (например, PageObjectWriter и TestClassWriter "
     "читают свойства сущности при генерации автотестов). Перечисления EntityKind, AttrType и "
     "ModifyType выступают управляющими признаками («↑●»): их значения определяют ветвление "
     "логики классификации сущностей и генерации тест-методов.")
figure("module_files_constantine.png", 2,
       "Структурная карта Константайна: обращения модулей и потоки данных")

body("Полный перечень модулей с указанием пакета, типа и назначения приведён в таблице 1.")
rows = []
for pkg, fn, cls, raw, _consts in CM.MODULES:
    rows.append([pkg, fn, TYPE_RU[MF.KIND[cls]], DESC.get(cls, "")])
table(1, "Модули программного средства",
      ["Пакет", "Модуль (файл .java)", "Тип", "Назначение"],
      [1.9, 4.6, 2.6, 7.4], rows)

doc.save(OUT)
print("Сохранено:", OUT)
print(f"Модулей: {n_mod} | пакетов: {n_pkg} | виды: {dict(kc)}")
missing = [cls for _p, _f2, cls, _r, _c in CM.MODULES if cls not in DESC]
print("Без описания:", missing if missing else "нет")
