# -*- coding: utf-8 -*-
"""Сборка отдельного документа «Модульная структура программного средства (модули и
константы)»: модуль = файл .java, внутри — его константы (все final-поля и enum-константы).
Стиль/нумерация — как в разделах 1.5 (Times New Roman, поля ГОСТ, подпись таблицы в две
строки, рисунок по центру). Самостоятельный документ — нумерация рисунка/таблицы с 1.
Хелперы оформления скопированы из build_section_153_docx.py."""
import os, sys, subprocess
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import const_model as CM

DIA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
ROOT = os.path.join(DIA, "..")
OUT = os.path.join(ROOT, "Модульная_структура_модули_константы.docx")
FONT = "Times New Roman"
USABLE_W = 16.5
MAX_H = 21.0
KIND_RU = {"class": "класс", "enum": "перечисление", "interface": "интерфейс", "record": "запись"}
PKG_RU = {"ui": "ui", "parser": "parser", "model": "model",
          "generator": "generator", "data": "data", "common": "common"}

doc = Document()
st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(14)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)
sec = doc.sections[0]
sec.left_margin = Cm(3); sec.right_margin = Cm(1.5); sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)


def _f(run, size=14, bold=False, italic=False):
    run.font.name = FONT; run.font.size = Pt(size); run.bold = bold; run.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def plural(n, one, few, many):
    n10, n100 = n % 10, n % 100
    if n10 == 1 and n100 != 11:
        return one
    if 2 <= n10 <= 4 and not 12 <= n100 <= 14:
        return few
    return many


def heading(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6); p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.5
    _f(p.add_run(text), 14, bold=False); return p


def body(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.75); p.paragraph_format.line_spacing = 1.5
    _f(p.add_run(text), 14); return p


def gap(n):
    for _ in range(n):
        p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.5
        _f(p.add_run(""), 14)


def figure(fname, num, caption):
    w, h = Image.open(os.path.join(DIA, fname)).size
    disp = USABLE_W
    if (h / w) * USABLE_W > MAX_H:
        disp = MAX_H * w / h
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5; p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(os.path.join(DIA, fname), width=Cm(disp))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.line_spacing = 1.5
    _f(c.add_run(f"Рисунок {num} – {caption}"), 14)
    gap(1)


def _shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexc); tcPr.append(sh)


def _ct(cell, text, bold=False):
    cell.text = ""; p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
    _f(p.add_run(text), 12, bold=bold)


def _ct_consts(cell, consts):
    """Список констант в ячейке: static final / enum-константы — прямым, final-поля
    экземпляра — курсивом. Пусто → «—»."""
    cell.text = ""; p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
    if not consts:
        _f(p.add_run("—"), 12); return
    for i, c in enumerate(consts):
        _f(p.add_run(c["name"]), 12, italic=not c["static"])
        if i < len(consts) - 1:
            _f(p.add_run(", "), 12)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true"); trPr.append(el)


def _no_split(row):
    trPr = row._tr.get_or_add_trPr(); el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true"); trPr.append(el)


def table(num, caption, headers, widths, modules):
    p1 = doc.add_paragraph(); p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.paragraph_format.line_spacing = 1.0; p1.paragraph_format.space_before = Pt(6)
    _f(p1.add_run(f"Таблица {num}"), 14)
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER; p2.paragraph_format.line_spacing = 1.0
    _f(p2.add_run(caption), 14)
    t = doc.add_table(rows=1 + len(modules), cols=len(headers))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.LEFT; t.autofit = False
    el = OxmlElement("w:tblLayout"); el.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(el)
    grid = t._tbl.find(qn("w:tblGrid"))          # фиксированный layout берёт ширины из tblGrid
    for gc, wd in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(round(wd * 567))))
    for j, h in enumerate(headers):
        _ct(t.rows[0].cells[j], h, bold=True); _shade(t.rows[0].cells[j], "D9D9D9")
    _repeat_header(t.rows[0])
    for i, (pkg, fn, cls, kind, consts) in enumerate(modules, 1):
        _no_split(t.rows[i])
        _ct(t.rows[i].cells[0], PKG_RU.get(pkg, pkg))
        _ct(t.rows[i].cells[1], fn)
        _ct(t.rows[i].cells[2], KIND_RU.get(kind, kind))
        _ct_consts(t.rows[i].cells[3], consts)
    for i in range(len(modules) + 1):
        for j, wd in enumerate(widths):
            t.rows[i].cells[j].width = Cm(wd)
    gap(2)


# ============================================================ содержимое
mods = CM.MODULES
n_mod = len(mods)
n_with = sum(1 for m in mods if m[4])
n_const = sum(len(m[4]) for m in mods)
n_enum = sum(1 for m in mods for c in m[4] if c["enum"])
n_static = sum(1 for m in mods for c in m[4] if c["static"] and not c["enum"])
n_inst = sum(1 for m in mods for c in m[4] if not c["static"])
n_pkg = len({m[0] for m in mods})

heading("Модульная структура программного средства (модули и константы)")
body("Информационная система AutoTestGenerator реализована на языке Java и организована по "
     "модульному принципу. В качестве модуля рассматривается отдельный файл исходного кода "
     f"(.java); всего система содержит {n_mod} {plural(n_mod, 'модуль', 'модуля', 'модулей')}, "
     f"распределённых по {n_pkg} пакетам (ui, parser, model, generator, data, common). "
     "Для каждого модуля на схеме показаны его константы.")
body("Под константой модуля понимается любое его поле, объявленное с модификатором final: "
     "именованная константа (static final), значение перечисления (enum-константа), а также "
     "неизменяемое поле экземпляра (final). Всего по модулям выявлено "
     f"{n_const} {plural(n_const, 'поле', 'поля', 'полей')} с модификатором final: из них "
     f"{n_static} — константы static final, {n_enum} — значения перечислений (enum-константы) "
     f"и {n_inst} — неизменяемые поля экземпляра. Константы имеются в {n_with} "
     f"{plural(n_with, 'модуле', 'модулях', 'модулях')} из {n_mod}; у остальных модулей "
     "собственные константы отсутствуют (помечено «нет констант»).")

body("Модульная структура системы — разбиение программного средства на пакеты и модули (.java) "
     "с перечислением констант каждого модуля — приведена на рисунке 1. Корневой узел "
     "соответствует программному средству в целом, пунктирные области — пакетам, прямоугольники — "
     "модулям; в верхней строке прямоугольника указано имя файла модуля, ниже — его константы. "
     "Константы static final и значения перечислений набраны прямым шрифтом, неизменяемые поля "
     "экземпляра — курсивом.")
figure("module_const_struct.png", 1, "Модульная структура программного средства: модули (файлы) и их константы")

body("Полный перечень модулей с указанием пакета, типа модуля и его констант приведён в "
     "таблице 1. В графе «Константы» значения перечислений и константы static final набраны "
     "прямым шрифтом, неизменяемые поля экземпляра (final) — курсивом; прочерк означает, что "
     "собственных констант у модуля нет.")
table(1, "Модули программного средства и их константы (поля final)",
      ["Пакет", "Модуль (файл .java)", "Тип", "Константы (поля final)"],
      [2.3, 5.3, 2.9, 6.0], mods)

doc.save(OUT)
print("Сохранено:", OUT)
print(f"Модулей: {n_mod} | с константами: {n_with} | всего констант: {n_const} "
      f"(static final {n_static}, enum {n_enum}, поля экземпляра {n_inst})")
