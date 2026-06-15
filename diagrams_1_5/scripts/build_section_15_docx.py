# -*- coding: utf-8 -*-
"""Сборка раздела 1.5.1–1.5.2 (Раздел_1.5_исправленный.docx) из реального кода.
Диаграммы (PNG) и таблицы строятся из одного источника (api.json), поэтому число
классов/полей/методов на рисунках и в таблицах совпадает. Нумерация рис./табл. с 100.
Формат: текст TNR 14, инт.1.5, по ширине, кр.строка 0.75; таблицы TNR 12, инт.1.0,
без отступа, по левому краю; рисунок по центру TNR 14 инт.1.5; зазор после таблицы
2 пустых абзаца (1.5), после рисунка — 1.
"""
import os, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import apimodel as M
import scenarios as S

DIA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
ROOT = os.path.join(DIA, "..")
ONLY = sys.argv[1] if len(sys.argv) > 1 else None   # имя пакета -> собрать ОДИН пакет (образец)
OUT = os.path.join(ROOT, f"Образец_1.5_{ONLY}.docx" if ONLY else "Раздел_1.5_исправленный.docx")
FONT = "Times New Roman"
USABLE_W = 16.5
MAX_H = 21.5

api = M.load_api(); idx = M.build_index(api)
PKG_TITLE = {"ui": "UI", "parser": "Parser", "model": "Model", "generator": "Generator",
             "data": "Data", "common": "Common"}

doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT; st.font.size = Pt(14)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)
sec = doc.sections[0]
sec.left_margin = Cm(3); sec.right_margin = Cm(1.5); sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)

fig_no = 99
tbl_no = 99


def _f(run, size=14, bold=False, italic=False, name=FONT):
    run.font.name = name; run.font.size = Pt(size); run.bold = bold; run.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def heading(text, size=14):
    # заголовки разделов — как обычный текст: 14 pt, без жирного, интервал 1.5
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6); p.paragraph_format.keep_with_next = True
    p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.line_spacing = 1.5
    _f(p.add_run(text), 14, bold=False); return p


def body(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.75); p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    _f(p.add_run(text), 14); return p


def gap(n):
    for _ in range(n):
        p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.space_after = Pt(0)
        _f(p.add_run(""), 14)


def next_fig():
    global fig_no; fig_no += 1; return fig_no


def next_tbl():
    global tbl_no; tbl_no += 1; return tbl_no


def figure(fname, num, caption):
    path = os.path.join(DIA, fname)
    w, h = Image.open(path).size
    disp_w = USABLE_W
    if (h / w) * USABLE_W > MAX_H:
        disp_w = MAX_H * w / h
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.keep_together = True
    p.add_run().add_picture(path, width=Cm(disp_w))
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0); cap.paragraph_format.line_spacing = 1.5
    _f(cap.add_run(f"Рисунок {num} – {caption}"), 14)
    gap(1)


def _shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexc); tcPr.append(sh)


def _ct(cell, text, bold=False):
    cell.text = ""; p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Cm(0)
    _f(p.add_run(text), 12, bold=bold)


def _fixed(t):
    el = OxmlElement("w:tblLayout"); el.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(el)


def table(num, caption, headers, rows, widths):
    # подпись: «Таблица N» — отдельная строка по правому краю; название — по центру
    p1 = doc.add_paragraph(); p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.paragraph_format.first_line_indent = Cm(0); p1.paragraph_format.line_spacing = 1.0
    p1.paragraph_format.space_before = Pt(6)
    _f(p1.add_run(f"Таблица {num}"), 14)
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.first_line_indent = Cm(0); p2.paragraph_format.line_spacing = 1.0
    p2.paragraph_format.space_after = Pt(2)
    _f(p2.add_run(caption), 14)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.LEFT; t.autofit = False
    _fixed(t)
    for j, h in enumerate(headers):
        _ct(t.rows[0].cells[j], h, bold=False); _shade(t.rows[0].cells[j], "D9D9D9")
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            _ct(t.rows[i].cells[j], val)
    for i in range(len(rows) + 1):
        for j, wd in enumerate(widths):
            t.rows[i].cells[j].width = Cm(wd)
    gap(2)


ROMAN = {"ui": 1, "parser": 2, "model": 3, "generator": 4, "data": 5, "common": 6}
FLOW_FIG = {"normal": "при нормальном ходе событий", "user": "при прерывании пользователем",
            "system": "при прерывании системой"}

MAINCONTROLLER_NOTE = (
    "Почти все методы класса MainController имеют закрытую видимость (-). Это объясняется тем, что "
    "MainController — контроллер графического интерфейса JavaFX, связанный с разметкой FXML. Обработчики "
    "событий (onParse, onGenerate, onRunTests и др.) и вспомогательные методы помечены аннотацией @FXML и "
    "вызываются загрузчиком FXMLLoader через механизм рефлексии, поэтому объявлены private. Публичным "
    "является только метод initialize() — он определён контрактом интерфейса Initializable и вызывается "
    "средой JavaFX при инициализации окна. Таким образом, класс не предоставляет открытого программного "
    "интерфейса, кроме инициализации, что является нормой для FXML-контроллеров.")


def package_section(pkg, X):
    title = PKG_TITLE[pkg]
    types = api[pkg]
    n = len(types)
    layer, name, desc = M.PKG_META[pkg]
    heading(f"1.5.2.{X} Разработка классов пакета «{title}»", 13)
    body(f"{desc} Пакет содержит {n} классов.")

    # .1 исходная диаграмма + таблица классов
    heading(f"1.5.2.{X}.1 Разработка исходной диаграммы классов пакета «{title}»", 13)
    f1 = next_fig()
    body(f"Исходная диаграмма классов пакета «{title}» приведена на рисунке {f1}. Классы соединены простыми "
         f"направленными связями; виды связей и кратности уточняются далее.")
    figure(f"cls_{pkg}_source.png", f1, f"Исходная диаграмма классов пакета «{title}»")
    t1 = next_tbl()
    body(f"Перечень классов пакета «{title}» приведён в таблице {t1}.")
    crows = [[r["qualified"], M.CLASS_RU.get(f"{pkg}.{r['qualified']}", "")] for r in types]
    for el in S.external_labels(pkg):
        ep = el.split("::")[0]
        crows.append([el, f"Внешний класс пакета {PKG_TITLE.get(ep, ep)}; описан в подразделе 1.5.2.{ROMAN[ep]}"])
    table(t1, f"Классы пакета «{title}»", ["Класс", "Назначение"], crows, [5.0, 11.5])

    # .2 диаграммы последовательности
    heading(f"1.5.2.{X}.2 Разработка диаграмм последовательности взаимодействия объектов классов пакета «{title}»", 13)
    flows = S.SC[pkg]["flows"]
    note = S.SC[pkg].get("note")
    intro_nums = [(fk, next_fig()) for fk in ("normal", "user", "system") if fk in flows]
    refs = ", ".join(f"рисунок {num} ({S.FLOW_TITLE[fk]})" for fk, num in intro_nums)
    body(f"Взаимодействие объектов классов пакета «{title}» показано диаграммами последовательности для "
         f"сценариев: {refs}. На всех трёх диаграммах присутствует один и тот же набор линий жизни; объекты, "
         f"не участвующие в конкретном сценарии, остаются без сообщений. Управляющее воздействие поступает из "
         f"края диаграммы (действующее лицо «Пользователь» не изображается), каждая активация завершается "
         f"возвратом управления, а прерывание показано знаком «✕». Классы других пакетов помечены «Пакет::Класс».")
    if note:
        body(note)
    for fk, num in intro_nums:
        figure(f"seq_{pkg}_{fk}.png", num,
               f"Диаграмма последовательности взаимодействия объектов классов пакета «{title}» {FLOW_FIG[fk]}")

    # .3 кооперация
    heading(f"1.5.2.{X}.3 Разработка диаграммы кооперации пакета «{title}»", 13)
    f3 = next_fig()
    objs = ", ".join(S.label_of(pkg, k) for k in S.used_object_keys(pkg))
    body(f"Диаграмма кооперации (рисунок {f3}) представляет те же объекты, что и диаграммы последовательности, "
         f"с нумерацией сообщений по сценариям (нормальный ход — 1, 2, …; прерывание пользователем — п1, п2, …; "
         f"прерывание системой — с1, с2, …). Участвующие объекты: {objs}.")
    figure(f"coop_{pkg}.png", f3, f"Диаграмма кооперации пакета «{title}»")

    # .4 уточнённая
    heading(f"1.5.2.{X}.4 Разработка уточнённой диаграммы классов пакета «{title}»", 13)
    f4 = next_fig()
    body(f"Уточнённая диаграмма классов пакета «{title}» (рисунок {f4}) фиксирует виды связей: обобщение (▷), "
         f"агрегацию (◇), композицию (◆) и ассоциацию (→), а также кратности на обоих концах связей.")
    figure(f"cls_{pkg}_refined.png", f4, f"Уточнённая диаграмма классов пакета «{title}»")

    # .5 детальная + таблицы полей/методов (геттеры/сеттеры опущены)
    heading(f"1.5.2.{X}.5 Разработка детальной диаграммы классов пакета «{title}»", 13)
    f5 = next_fig()
    body(f"Детальная диаграмма классов пакета «{title}» (рисунок {f5}) содержит поля и методы каждого класса с "
         f"указанием признака видимости (+ открытый, - закрытый, # защищённый, ~ пакетный). Тривиальные геттеры "
         f"и сеттеры полей для наглядности не приводятся (ни на диаграмме, ни в таблицах). Состав полей и методов "
         f"классов приведён в таблицах ниже.")
    figure(f"cls_{pkg}_detailed.png", f5, f"Детальная диаграмма классов пакета «{title}»")
    if pkg == "ui":
        body(MAINCONTROLLER_NOTE)
    for r in types:
        pkg_qual = f"{pkg}.{r['qualified']}"
        if r["fields"]:
            tnum = next_tbl()
            body(f"Поля класса {r['qualified']} приведены в таблице {tnum}.")
            frows = [[f["name"], M.xref(f["type"], idx, pkg), M.describe_field(f)] for f in r["fields"]]
            table(tnum, f"Поля класса {r['qualified']}", ["Название", "Тип", "Описание"],
                  frows, [3.8, 4.2, 8.5])
        vmeth = M.visible_methods(r)
        if vmeth:
            tnum = next_tbl()
            body(f"Методы класса {r['qualified']} приведены в таблице {tnum}.")
            mrows = []
            for m in vmeth:
                params = ", ".join(f"{p['name']}: {M.xref(p['type'], idx, pkg)}" for p in m["params"]) or "—"
                ret = "—" if m["ctor"] else M.xref(m["returns"], idx, pkg)
                mrows.append([m["name"], params, ret, M.describe_method(pkg_qual, m)])
            table(tnum, f"Методы класса {r['qualified']}",
                  ["Название метода", "Параметры", "Возвращаемое значение", "Описание"],
                  mrows, [3.6, 4.4, 3.0, 5.5])
    for el in S.external_labels(pkg):
        ep = el.split("::")[0]
        body(f"Класс {el} используется классами пакета «{title}», но относится к пакету "
             f"{PKG_TITLE.get(ep, ep)}; его поля и методы приведены в подразделе 1.5.2.{ROMAN[ep]} и здесь не "
             f"дублируются.")


# ====================================================================== содержимое
if ONLY:
    heading(f"1.5.2.{ROMAN[ONLY]} Разработка классов пакета «{PKG_TITLE[ONLY]}» (образец)", 14)
    body("Образец одного пакета для проверки оформления. Нумерация рисунков и таблиц начата с № 100 и "
         "подлежит сквозной перенумерации. После согласования будет собран весь раздел 1.5.")
    package_section(ONLY, ROMAN[ONLY])
    doc.save(OUT)
    print("Сохранено:", OUT)
    print(f"Рисунков: {fig_no - 99} (100..{fig_no}) | Таблиц: {tbl_no - 99} (100..{tbl_no})")
    sys.exit(0)

heading("1.5 Проектирование программного обеспечения", 16)
body("Проектирование программного обеспечения при объектном подходе включает разработку структуры системы "
     "(деление на пакеты), уточнение отношений между классами, проектирование взаимодействия объектов и "
     "детальное проектирование классов. Раздел построен строго по реальному коду системы: состав классов, "
     "полей и методов на диаграммах классов и в таблицах совпадает с реализацией. Нумерация рисунков и таблиц "
     "в данном разделе начата с № 100 и подлежит сквозной перенумерации при включении в работу.")

# ---------------- 1.5.1 ----------------
heading("1.5.1 Проектирование структуры системы и построение диаграмм пакетов", 14)
total = sum(len(api[p]) for p in M.PKG_ORDER)
fpkg = next_fig()
body(f"Программная система разделена на шесть пакетов единого назначения, сгруппированных по слоям. Всего "
     f"система содержит {total} классов (включая вложенные). Структура пакетов и зависимости между ними по "
     f"направлению вызовов приведены на рисунке {fpkg}.")
figure("pkg_diagram.png", fpkg, "Диаграмма пакетов информационной системы")
tpkg = next_tbl()
body(f"Назначение пакетов, их принадлежность слоям и число классов приведены в таблице {tpkg}.")
rows = []
for p in M.PKG_ORDER:
    layer, name, desc = M.PKG_META[p]
    rows.append([name, layer, str(len(api[p])), desc])
table(tpkg, "Пакеты программной системы", ["Пакет", "Слой", "Классов", "Назначение"],
      rows, [2.3, 3.2, 1.8, 9.2])

# ---------------- 1.5.2 ----------------
heading("1.5.2 Проектирование классов в пакетах", 14)
body("Для каждого пакета построены исходная диаграмма классов, диаграммы последовательности взаимодействия "
     "объектов (нормальный ход, прерывание пользователем, прерывание системой), диаграмма кооперации, "
     "уточнённая и детальная диаграммы классов. На всех диаграммах классов пакета показано одинаковое число "
     "классов, равное числу классов пакета в коде; на диаграммах взаимодействия управление поступает из края "
     "(действующее лицо не показывается), а классы других пакетов помечаются в формате «Пакет::Класс».")

for pkg in M.PKG_ORDER:
    package_section(pkg, ROMAN[pkg])

doc.save(OUT)
print("Сохранено:", OUT)
print(f"Рисунков: {fig_no - 99} (100..{fig_no}) | Таблиц: {tbl_no - 99} (100..{tbl_no})")
