# -*- coding: utf-8 -*-
"""Сборка отдельного раздела «Построение модульной структуры» на 18 модулей
(старая схема) в оформлении ВКР Сафиуллина: рисунок модульной структуры,
краткая спецификация (4 столбца), подробная спецификация по каждому модулю
(3 столбца), карта Константайна, связность (2 столбца) и сцепление (2 столбца).
Диаграммы/таблицы компонентов НЕ включаются. Вывод: Модульная_структура_18модулей.docx."""
import os, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import module_model_s18 as MM

DIA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
ROOT = os.path.join(DIA, "..")
OUT = os.path.join(ROOT, "Модульная_структура_18модулей.docx")
FONT = "Times New Roman"
USABLE_W = 16.5; MAX_H = 21.0
fig_no = 142; tbl_no = 179

doc = Document()
st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(14)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)
sec = doc.sections[0]
sec.left_margin = Cm(3); sec.right_margin = Cm(1.5); sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)


def _f(run, size=14, bold=False, italic=False):
    run.font.name = FONT; run.font.size = Pt(size); run.bold = bold; run.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def heading(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6); p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.5
    _f(p.add_run(text), 14, bold=False); return p


def body(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.75); p.paragraph_format.line_spacing = 1.5
    _f(p.add_run(text), 14); return p


def gap(n):
    for _ in range(n):
        p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.5
        _f(p.add_run(""), 14)


def next_fig():
    global fig_no; fig_no += 1; return fig_no


def next_tbl():
    global tbl_no; tbl_no += 1; return tbl_no


def figure(fname, num, caption):
    w, h = Image.open(os.path.join(DIA, fname)).size
    disp = USABLE_W
    if (h / w) * USABLE_W > MAX_H:
        disp = MAX_H * w / h
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5; p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(os.path.join(DIA, fname), width=Cm(disp))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.line_spacing = 1.5
    _f(c.add_run(f"Рисунок {num} – {caption}"), 14)
    gap(1)


def _shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexc); tcPr.append(sh)


def _ct(cell, text):
    cell.text = ""; p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
    _f(p.add_run(text), 12)


def table(num, caption, headers, rows, widths):
    p1 = doc.add_paragraph(); p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.paragraph_format.line_spacing = 1.0; p1.paragraph_format.space_before = Pt(6)
    p1.paragraph_format.keep_with_next = True
    _f(p1.add_run(f"Таблица {num}"), 14)
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER; p2.paragraph_format.line_spacing = 1.0
    p2.paragraph_format.keep_with_next = True
    _f(p2.add_run(caption), 14)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.LEFT; t.autofit = False
    el = OxmlElement("w:tblLayout"); el.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(el)
    for j, h in enumerate(headers):
        _ct(t.rows[0].cells[j], h); _shade(t.rows[0].cells[j], "D9D9D9")
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            _ct(t.rows[i].cells[j], v)
    for i in range(len(rows) + 1):
        for j, wd in enumerate(widths):
            t.rows[i].cells[j].width = Cm(wd)
    gap(2)


# ============================================================ содержимое
N = len(MM.MODULES18)
heading("1.5.3 Построение модульной структуры")
body("Модульная структура – диаграмма программного обеспечения, в которой приложение разбивается на "
     "независимые компоненты – модули. Каждый модуль выполняет определённую функцию, имеет собственную логику "
     "и интерфейс для взаимодействия с другими модулями.")
f1 = next_fig()
t_short = tbl_no + 1
body(f"Структура программы состоит из {N} модулей (рисунок {f1}), её краткая спецификация приведена в "
     f"таблице {t_short}. Центральным элементом системы является модуль MainController, который координирует "
     "разбор XML-метамодели, генерацию проекта автотестов, их запуск и формирование отчётов.")
names = ", ".join(m[1] for m in MM.MODULES18)
body("В состав системы входят следующие функциональные модули и повторно используемые библиотеки: " + names + ".")
figure("mod_struct_18.png", f1, "Модульная структура информационной системы")

t1 = next_tbl()
table(t1, "Краткая спецификация модулей системы",
      ["Наименование", "Входные данные", "Выходные данные", "Описание"],
      [[m[1], m[2], m[3], m[4]] for m in MM.MODULES18], [3.2, 3.6, 3.6, 6.1])

# подробная спецификация по каждому модулю
t_first = tbl_no + 1
t_last = tbl_no + N
body(f"Подробная спецификация модулей системы (перечень функций и процедур каждого модуля с параметрами и "
     f"назначением) приведена в таблицах {t_first}–{t_last}.")
for m in MM.MODULES18:
    key, name = m[0], m[1]
    tn = next_tbl()
    table(tn, f"Подробная спецификация модуля «{name}»",
          ["Название функции или процедуры", "Параметры", "Описание"],
          [list(r) for r in MM.METHODS[key]], [5.3, 4.2, 7.0])

# карта Константайна
f2 = next_fig()
body(f"Для анализа структуры системы и оценки взаимодействия её модулей применяется карта Константайна "
     f"(рисунок {f2}). На ней вершинами выступают модули системы (прямоугольник), повторно используемые "
     "библиотеки (двойная рамка) и общие области хранения данных (овал), а дугами – информационные связи "
     "и потоки данных. Все обращения показаны сплошными стрелками; у связей подписаны куплеты с "
     "направлением (↓ – передача в вызываемый модуль, ↑ – возврат) и видом связи (○ – по данным, "
     "● – по управлению).")
figure("constantine_18.png", f2, "Структурная карта Константайна информационной системы")

# связность и сцепление
body("Качество построенной модульной структуры оценивается показателями связности и сцепления. Связность "
     "модуля – это мера зависимости его частей; чем выше связность, тем лучше результат проектирования. "
     "Для рассматриваемой системы большинство модулей обладают функциональной связностью, парсеры-помощники "
     "и библиотеки – логической и информационной, а модули-оркестраторы – последовательной связностью.")
t_coh = next_tbl()
table(t_coh, "Связность модулей", ["Наименование", "Тип связи"],
      [[m[1], m[5]] for m in MM.MODULES18], [9.0, 7.5])

body("Сцепление определяет уровень зависимости между отдельными модулями системы. Наиболее рациональным "
     "считается сцепление по данным. В разрабатываемой системе взаимодействие осуществляется преимущественно "
     "по образцу (передаются составные объекты модели) и по данным; обращения к повторно используемым "
     "библиотекам (StaxUtils, Transliterator, JavaFileWriter, XmlNamespaces) – сцепление по данным "
     "(передаются параметры и результат, без общего изменяемого состояния), а возврат вида сущности "
     "классификатором – сцепление по управлению. Оценка сцепления модулей приведена в таблице ниже.")
t_cou = next_tbl()
table(t_cou, "Сцепление модулей", ["Вызывающий модуль", "Вид сцепления"],
      [[m[1], MM.COUPLING18[m[0]]] for m in MM.MODULES18], [9.0, 7.5])

doc.save(OUT)
print(f"Сохранено: {OUT}")
print(f"Рисунков: 2 ({f1}, {f2}) | Таблиц: {tbl_no - 179} (180..{tbl_no})")
