# -*- coding: utf-8 -*-
"""Сборка Приложения 3 (Приложение_3_Спецификация.docx) — спецификации на
разработанную программную документацию и программное обеспечение.

Таблица перечисляет все физические компоненты системы AutoTestGenerator
(файл сборки pom.xml, ресурс main.fxml и 42 класса Java по пакетам).

Формат: Times New Roman 12 pt, межстрочный интервал 1.0, поля 3/1.5/2/2 см,
таблица — стиль Table Grid с заливкой шапки.
Запуск:  python3 prilozhenie_3/build_prilozhenie_3_docx.py
"""
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, "..")
OUT = os.path.join(ROOT, "Приложение_3_Спецификация.docx")
FONT = "Times New Roman"
USABLE_CM = 16.5
COLS = (4.5, 8.5, 3.5)  # Обозначение / Описание / Примечание  (сумма = 16.5)

# --------------------------------------------------------------------------- #
#  Содержимое таблицы                                                          #
# --------------------------------------------------------------------------- #
DOC_TITLE = (
    "Расчётно-пояснительная записка выпускной квалификационной работы на тему "
    "«Разработка информационной системы генерации автотестов для web-приложения "
    "на основе xml-файла»"
)

DOCUMENTATION = [
    ("Расчётно-пояснительная записка VKR_Safiullin.docx", DOC_TITLE,
     "Расчётно-пояснительная записка в формате docx"),
    ("Приложение 2", "Текст программы",
     "Содержит исходный программный код разработанной информационной системы."),
    ("Приложение 4", "Руководство пользователя",
     "Содержит инструкции по установке, настройке и эксплуатации разработанной системы."),
]

# Каждый кортеж: (Обозначение, Описание, Примечание)
COMPONENTS = [
    # --- Сборка и ресурсы ---
    ("pom.xml",
     "Файл сборки проекта Maven: задаёт координаты артефакта, зависимости "
     "(JavaFX, SQLite JDBC) и плагины сборки.",
     "Управляет сборкой проекта и подключением библиотек."),
    ("main.fxml",
     "Декларативное FXML-описание главного окна: поля ввода параметров, кнопки, "
     "дерево сущностей, таблица результатов и область лога.",
     "Загружается классом App, связан с MainController."),

    # --- Пакет ui (пользовательский интерфейс) ---
    ("App.java",
     "Точка входа приложения, инициализирует JavaFX и открывает главное окно программы.",
     "Запускается первым при старте приложения."),
    ("Launcher.java",
     "Вспомогательная точка входа для упакованного приложения (jpackage); "
     "передаёт управление классу App.",
     "Используется при запуске собранного дистрибутива."),
    ("MainController.java",
     "Контроллер главного окна, обрабатывает действия пользователя и координирует "
     "работу остальных модулей.",
     "Центральный модуль пользовательского интерфейса."),

    # --- Пакет parser (разбор XML) ---
    ("XmlModelParser.java",
     "Выполняет потоковый разбор XML-метамодели и формирует объектную модель приложения.",
     "Использует технологию StAX."),
    ("EntityParser.java",
     "Разбирает XML-блок <Object> в объект сущности, включая свойства и ассоциации.",
     "Вызывается из XmlModelParser."),
    ("PropertyGroupParser.java",
     "Разбирает XML-блок <Properties> в группу свойств сущности с полями и операцией.",
     "Вызывается при разборе сущности."),
    ("SearchParser.java",
     "Разбирает XML-блок <Searches> в список объектов поиска с параметрами и "
     "описанием грида результатов.",
     "Вызывается из XmlModelParser."),
    ("StaxUtils.java",
     "Набор статических вспомогательных методов для парсеров: доступ к атрибутам, "
     "безопасный разбор чисел, перемотка курсора.",
     "Общая утилита пакета parser."),
    ("XmlNamespaces.java",
     "Константы XML-пространств имён формата метаданных E3Core.",
     "Используется парсерами."),

    # --- Пакет model (объектная модель) ---
    ("AppModel.java",
     "Корневой объект разобранной модели: содержит списки сущностей и поисков.",
     "Результат работы парсера."),
    ("EntityObject.java",
     "Описывает сущность модели: имя, ключ, ассоциации и группы свойств.",
     "Элемент объектной модели."),
    ("PropertyGroup.java",
     "Группа свойств сущности (форма, вкладка или грид) с возможной операцией.",
     "Элемент объектной модели."),
    ("Property.java",
     "Описывает поле (атрибут) сущности: тип, маску, обязательность, порядок.",
     "Элемент объектной модели."),
    ("Operation.java",
     "Описывает операцию над данными (создание, изменение, удаление) с параметрами "
     "и модификаторами.",
     "Элемент объектной модели."),
    ("OperationParam.java",
     "Параметр операции над данными.",
     "Элемент объектной модели."),
    ("Modifier.java",
     "Модификатор операции, задающий тип изменения данных.",
     "Элемент объектной модели."),
    ("ModifyType.java",
     "Перечисление типов изменения данных (вставка, обновление, удаление и др.).",
     "Перечисление."),
    ("Association.java",
     "Описывает ассоциацию (связь) между сущностями модели.",
     "Элемент объектной модели."),
    ("AttrType.java",
     "Перечисление типов атрибутов (строка, число, дата, дата-время).",
     "Перечисление."),
    ("Search.java",
     "Описывает параметрический поиск сущности.",
     "Элемент объектной модели."),
    ("SearchParam.java",
     "Параметр поиска: имя, заголовок, тип значения, маска.",
     "Элемент объектной модели."),
    ("SearchResult.java",
     "Описание грида результатов поиска.",
     "Элемент объектной модели."),
    ("SearchResultProperty.java",
     "Колонка грида результатов поиска.",
     "Элемент объектной модели."),
    ("EntityKind.java",
     "Перечисление категорий сущности (главная, дочерняя, справочник).",
     "Перечисление."),
    ("EntityClassifier.java",
     "Классифицирует сущности модели на главные, дочерние и справочники для выбора "
     "стратегии генерации тестов.",
     "Определяет, какие тесты создавать."),
    ("TestCaseResult.java",
     "Результат отдельного тест-метода: статус, длительность, скриншоты, перехваченный вывод.",
     "Используется в отчётах."),
    ("TestRunResult.java",
     "Агрегированный результат прогона набора тестов (всего, успешно, ошибки, пропущено).",
     "Используется в отчётах."),

    # --- Пакет generator (генерация и запуск тестов) ---
    ("TestGenerator.java",
     "Создаёт структуру генерируемого тест-проекта и формирует Page Object'ы и "
     "тест-классы для каждой сущности.",
     "Координирует генерацию тестов."),
    ("TestConfig.java",
     "Хранит параметры генерации: URL, учётные данные, каталог вывода, тип сайта, "
     "уровень тестов.",
     "Передаётся в генератор."),
    ("TestClassWriter.java",
     "Генерирует JUnit 5 тест-классы для каждой сущности: наличие полей, валидация, "
     "CRUD-операции, поиск.",
     "Формирует исходный код тестов."),
    ("PageObjectWriter.java",
     "Генерирует классы Page Object (Selenium) с элементами форм и методами работы с ними.",
     "Формирует исходный код тестов."),
    ("TestDataFactory.java",
     "Генерирует корректные тестовые значения по типу поля, маске и стереотипу.",
     "Используется при генерации тестов."),
    ("TestRunner.java",
     "Запускает сгенерированный проект тестов через Maven Surefire и разбирает "
     "XML-отчёты в результат прогона.",
     "Выполняет прогон тестов."),
    ("RunReportWriter.java",
     "Формирует самостоятельный HTML- и CSV-отчёт о прогоне: скриншоты, перехваченный "
     "вывод, параметры поиска, время по шагам.",
     "Формирует отчёт о прогоне."),

    # --- Пакет data (хранение отчётов) ---
    ("ReportDao.java",
     "Фасад над базой отчётов SQLite: сохранение прогонов и загрузка истории.",
     "Объединяет DAO-классы пакета."),
    ("DatabaseConnection.java",
     "Управляет JDBC-подключением к базе отчётов SQLite.",
     "Единственный источник JDBC-URL."),
    ("SchemaInitializer.java",
     "Создаёт схему базы данных (таблицы) при первом запуске.",
     "Выполняет DDL-операции."),
    ("TestRunDao.java",
     "DAO для таблицы прогонов test_run: добавление записи и выборка списка прогонов.",
     "Доступ к данным."),
    ("TestCaseDao.java",
     "DAO для таблицы тест-кейсов test_case: пакетная вставка и выборка по прогону.",
     "Доступ к данным."),
    ("autotestgen.db",
     "Локальная база данных SQLite с историей прогонов: таблица test_run (сведения "
     "о запусках) и таблица test_case (результаты отдельных тестов).",
     "Создаётся автоматически при первом запуске; схема задаётся классом SchemaInitializer."),

    # --- Пакет common (общие утилиты) ---
    ("JavaFileWriter.java",
     "Утилита форматированной записи Java-файлов при генерации исходного кода тестов.",
     "Вспомогательный модуль."),
    ("Transliterator.java",
     "Преобразует кириллические имена в корректные Java-идентификаторы.",
     "Вспомогательный модуль."),
    ("ParserException.java",
     "Прикладное исключение, выбрасываемое при ошибках разбора XML.",
     "Вспомогательный модуль."),
]

# --------------------------------------------------------------------------- #
#  Вспомогательные функции форматирования                                      #
# --------------------------------------------------------------------------- #
doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.line_spacing = 1.0
st.paragraph_format.space_after = Pt(0)
sec = doc.sections[0]
sec.left_margin = Cm(3)
sec.right_margin = Cm(1.5)
sec.top_margin = Cm(2)
sec.bottom_margin = Cm(2)


def _f(run, size=14, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def _shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexc)
    tcPr.append(sh)


def _ct(cell, text, bold=False, size=12, align=None, valign="center"):
    cell.text = ""
    cell.vertical_alignment = {"center": 1, "top": 0}.get(valign, 1)
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(0)
    _f(p.add_run(text), size, bold=bold)


def _set_widths(row):
    for j, w in enumerate(COLS):
        row.cells[j].width = Cm(w)


def _span_row(table, text):
    """Строка-разделитель раздела, объединённая на все три колонки."""
    row = table.add_row()
    merged = row.cells[0].merge(row.cells[1]).merge(row.cells[2])
    _ct(merged, text, bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    _shade(merged, "E8E8E8")


def _data_row(table, values):
    row = table.add_row()
    _set_widths(row)
    _ct(row.cells[0], values[0], size=12, valign="top")
    _ct(row.cells[1], values[1], size=12, valign="top")
    _ct(row.cells[2], values[2], size=12, valign="top")


def _repeat_header(row):
    """Помечает строку как повторяющуюся шапку на каждой странице."""
    trPr = row._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)


# --------------------------------------------------------------------------- #
#  Заголовок приложения                                                        #
# --------------------------------------------------------------------------- #
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
_f(p.add_run("Приложение 3"), 12, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(6)
_f(p.add_run("Спецификация"), 12, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.first_line_indent = Cm(1.25)
_f(p.add_run(
    "Настоящее приложение содержит спецификацию на разработанную программную "
    "документацию и программное обеспечение. Спецификация приведена в табл. П3.1."
), 12)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p.paragraph_format.space_before = Pt(8)
_f(p.add_run("Таблица П3.1"), 12)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
_f(p.add_run(
    "Спецификация на разработанную программную документацию и программное обеспечение"
), 12)

# --------------------------------------------------------------------------- #
#  Таблица                                                                     #
# --------------------------------------------------------------------------- #
table = doc.add_table(rows=2, cols=3)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER

# Шапка
hdr = table.rows[0]
_set_widths(hdr)
for j, h in enumerate(("Обозначение", "Описание", "Примечание")):
    _ct(hdr.cells[j], h, bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    _shade(hdr.cells[j], "D9D9D9")
_repeat_header(hdr)

# Нумерация колонок (1 / 2 / 3)
numr = table.rows[1]
_set_widths(numr)
for j, n in enumerate(("1", "2", "3")):
    _ct(numr.cells[j], n, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
_repeat_header(numr)

# Раздел «Документация»
_span_row(table, "Документация")
for vals in DOCUMENTATION:
    _data_row(table, vals)

# Раздел «Компоненты»
_span_row(table, "Компоненты")
for vals in COMPONENTS:
    _data_row(table, vals)

doc.save(OUT)
print("Сохранено:", OUT)
print("Строк документации:", len(DOCUMENTATION), "| компонентов:", len(COMPONENTS))
