# -*- coding: utf-8 -*-
"""Отдельный docx-файл с исправленным разделом 1.4.6 «Построение диаграммы
переходов состояний»: корректная диаграмма (по коду MainController) + текст,
согласованный с её состояниями. Файл Столбов_ВКР.docx не изменяется —
этот раздел можно вставить в диплом вместо текущего 1.4.6.

Форматирование: Times New Roman 14, межстрочный 1,5, красная строка 0,75 см;
рисунок и подпись — по центру."""
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from PIL import Image

DIA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
OUT = os.path.join(DIA, "..", "Раздел_1.4.6_исправленный.docx")
FONT = "Times New Roman"

doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT; st.font.size = Pt(14)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)
sec = doc.sections[0]
sec.left_margin = Cm(3); sec.right_margin = Cm(1.5); sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)


def _f(run, size=14, bold=False):
    run.font.name = FONT; run.font.size = Pt(size); run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def heading(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6); p.paragraph_format.keep_with_next = True
    _f(p.add_run(text), 14, bold=True); return p


def body(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.75)
    _f(p.add_run(text), 14); return p


def bullet(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(0.75); p.paragraph_format.first_line_indent = Cm(-0.5)
    _f(p.add_run("– " + text), 14); return p


def figure(fname, caption, width_cm):
    path = os.path.join(DIA, fname)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.keep_together = True
    p.add_run().add_picture(path, width=Cm(width_cm))
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10); cap.paragraph_format.keep_with_next = False
    _f(cap.add_run(caption), 14)


# ===================== Раздел 1.4.6 =====================
heading("1.4.6 Построение диаграммы переходов состояний")

body("Диаграмма переходов состояний (State Diagram) — графическая модель, описывающая поведение "
     "объекта или системы в ответ на события и последовательность изменений его состояний.")
body("Диаграмма состоит из нескольких ключевых элементов:")
bullet("Состояния (States). Основные состояния объекта, отображаемые в виде прямоугольников с "
       "закруглёнными углами. Каждое состояние характеризует определённую стадию жизненного цикла объекта.")
bullet("Переходы (Transitions). Стрелки, показывающие изменения состояния объекта. Переходы могут "
       "быть вызваны событиями или условиями, а сама стрелка подписывается именем события.")
bullet("События (Events). Внешние или внутренние события, которые инициируют переходы между состояниями.")
bullet("Начальное и конечное состояния (Initial and Final States). Начальное состояние обозначает "
       "стартовое положение объекта, конечное — момент завершения его жизненного цикла. Диаграмма "
       "переходов состояний системы показана на рис. 12.")

body("После запуска система переходит в состояние «Ожидание ввода», в котором пользователь задаёт "
     "путь к XML-метамодели, каталог сохранения проекта, адрес тестируемого сайта, учётные данные, тип "
     "сайта и наименование подсистемы. По команде «Разобрать метаданные» запускается потоковый "
     "StAX-парсер; если файл не указан, недоступен или имеет некорректную структуру, система выводит "
     "сообщение об ошибке и остаётся в состоянии ожидания ввода. После успешного разбора выполняется "
     "классификация сущностей на основные, дочерние и справочники, и система переходит в состояние "
     "«Метаданные разобраны», в котором становится доступной команда генерации. По команде "
     "«Сгенерировать автотесты» создаётся структура Maven-проекта (файл pom.xml, общие классы "
     "SharedDriver, BaseTest, TestData), а для каждой основной сущности — Page Object и тестовый класс "
     "JUnit 5; система переходит в состояние «Тесты сгенерированы». Из него по команде «Запустить "
     "автотесты» (или «Запустить выбранные») система переходит в состояние «Выполнение тестов» и "
     "выполняет mvn test с потоковым выводом хода прогона. По завершении прогона результаты "
     "отображаются, сведения о прогоне сохраняются в базе SQLite, формируются HTML-отчёт и "
     "CSV-выгрузка — система переходит в состояние «Результаты получены»; при ошибке запуска выводится "
     "сообщение, и система возвращается в состояние «Тесты сгенерированы», сохраняя возможность "
     "повторного запуска. В любой момент по команде «Просмотреть историю» доступен просмотр ранее "
     "сохранённых прогонов, а по команде выхода работа приложения завершается.")

figure("ris_12_state_transition.png", "Рис. 12. Диаграмма переходов состояний", 11.0)

doc.save(OUT)
print("Сохранено:", OUT)
