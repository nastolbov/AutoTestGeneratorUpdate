#!/usr/bin/env python3
# Приложение 2 «Текст программы» для ВКР (стиль РПЗ Столбова).
# Код — в ДВЕ КОЛОНКИ (через невидимую таблицу), Times New Roman 8pt, одинарный.
# Вступление/подпись — во всю ширину, 14pt, интервал 1.5, красная строка.
# Длинные листинги разбиваются по страницам: сверху справа «Продолжение рис. П2.N».
# Между рисунками — ровно 1 пустой абзац (Enter) 14pt / 1.5, без разрыва страницы.
import os
from docx import Document
from docx.shared import Pt, Emu, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = "/home/user/AutoTestGeneratorUpdate"
OUT = os.path.join(ROOT, "docs", "Приложение_2_Текст_программы.docx")

PACKAGES = ["ui", "parser", "model", "generator", "data", "common"]
JAVA_BASE = os.path.join(ROOT, "src/main/java/ru/autotestgen")
FXML = os.path.join(ROOT, "src/main/resources/fxml/main.fxml")

# --- параметры верстки ---
WRAP_WIDTH  = 58          # символов в строке кода (узкая колонка → перенос короче)
FIRST_BLOCK = 116         # строк кода на 1-й странице листинга (≈58 на колонку, под вступление)
CONT_BLOCK  = 128         # строк кода на странице-продолжении (≈64 на колонку)
RED_LINE = Emu(269875)    # красная строка как в РПЗ
COL_W = Inches(3.15)      # ширина одной колонки кода

# страница как в РПЗ (А4 с теми же полями) — чтобы разбивка совпала при вставке
PAGE_W, PAGE_H = Emu(7560310), Emu(10692130)
M_L, M_R, M_T, M_B = Emu(1080135), Emu(540385), Emu(720090), Emu(720090)

def files_for(pkg):
    d = os.path.join(JAVA_BASE, pkg)
    out = [os.path.join(d, f) for f in sorted(os.listdir(d)) if f.endswith(".java")]
    if pkg == "ui":
        out.append(FXML)
    return out

all_files = []
for pkg in PACKAGES:
    all_files.extend(files_for(pkg))

def wrap_code(line):
    line = line.replace("\t", "    ").rstrip("\n")
    if len(line) <= WRAP_WIDTH:
        return [line]
    indent = len(line) - len(line.lstrip(" "))
    cont = " " * min(indent + 4, 8)
    out, s, first = [], line, True
    while len(s) > WRAP_WIDTH:
        lo = (indent if first else len(cont)) + 1
        bp = s.rfind(" ", lo, WRAP_WIDTH + 1)
        if bp < lo or bp - len(cont) < 10:
            bp = WRAP_WIDTH
        out.append(s[:bp])
        s = cont + s[bp:].lstrip(" ")
        first = False
    out.append(s)
    return out

def read_code_lines(path):
    with open(path, encoding="utf-8") as f:
        raw = f.read().split("\n")
    if raw and raw[-1] == "":
        raw = raw[:-1]
    lines = []
    for ln in raw:
        lines.extend(wrap_code(ln))
    return lines

def split_blocks(lines):
    blocks, i = [], 0
    blocks.append(lines[i:i + FIRST_BLOCK]); i += FIRST_BLOCK
    while i < len(lines):
        blocks.append(lines[i:i + CONT_BLOCK]); i += CONT_BLOCK
    return [b for b in blocks if b]

# ====================== ДОКУМЕНТ ======================
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = PAGE_W, PAGE_H
sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = M_L, M_R, M_T, M_B

st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(12)
st.paragraph_format.line_spacing = 1.0
st.paragraph_format.space_after = Pt(0)

def big_par(text, align, indent=False, page_break=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(0)
    if indent:
        pf.first_line_indent = RED_LINE
    if page_break:
        pf.page_break_before = True
    p.alignment = align
    r = p.add_run(text); r.font.name = "Times New Roman"; r.font.size = Pt(14)
    return p

def enter_14():
    """Пустой абзац-разделитель между рисунками: 14pt, интервал 1.5."""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(""); r.font.name = "Times New Roman"; r.font.size = Pt(14)
    return p

def fill_code(cell, lines):
    """Заполняет ячейку кодом: TNR 8pt, одинарный, влево, без отступа."""
    p = cell.paragraphs[0]
    pf = p.paragraph_format
    pf.line_spacing = 1.0; pf.space_after = Pt(0); pf.space_before = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run()
    run.font.name = "Times New Roman"; run.font.size = Pt(8)
    rpr = run._element.get_or_add_rPr(); rf = rpr.get_or_add_rFonts()
    rf.set(qn("w:ascii"), "Times New Roman"); rf.set(qn("w:hAnsi"), "Times New Roman")
    for i, ln in enumerate(lines):
        if i > 0:
            run.add_break()
        run.add_text(ln if ln != "" else " ")

def set_cell_margins(cell, l=40, r=40, t=0, b=0):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for tag, val in (("top", t), ("start", l), ("bottom", b), ("end", r)):
        e = OxmlElement("w:" + tag); e.set(qn("w:w"), str(val)); e.set(qn("w:type"), "dxa")
        m.append(e)
    tcPr.append(m)

def code_two_columns(lines):
    """Невидимая таблица 1×2: левая колонка — первая половина строк, правая — вторая."""
    half = (len(lines) + 1) // 2
    left, right = lines[:half], lines[half:]
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.allow_autofit = False
    # строку не разрывать между страницами (разбивку контролируем сами)
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))
    for idx, content in enumerate((left, right)):
        c = table.cell(0, idx)
        c.width = COL_W
        set_cell_margins(c)
        if content:
            fill_code(c, content)
    return table

def kind_word(fname):
    return "файла" if fname.endswith(".fxml") else "класса"

# Заголовок приложения
big_par("Приложение 2", WD_ALIGN_PARAGRAPH.RIGHT)
big_par("Текст программы", WD_ALIGN_PARAGRAPH.CENTER, indent=True)
enter_14()

for idx, path in enumerate(all_files, start=1):
    fname = os.path.basename(path)
    num = "П2.%d" % idx
    blocks = split_blocks(read_code_lines(path))
    big_par("На рис. %s представлен текст %s %s." % (num, kind_word(fname), fname),
            WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True)
    for bi, block in enumerate(blocks):
        if bi > 0:
            big_par("Продолжение рис. %s" % num, WD_ALIGN_PARAGRAPH.RIGHT, page_break=True)
        code_two_columns(block)
    big_par("Рис. %s. Текст %s %s" % (num, kind_word(fname), fname), WD_ALIGN_PARAGRAPH.CENTER)
    enter_14()   # ровно 1 Enter (14pt/1.5) между рисунками, без разрыва страницы

doc.save(OUT)
print("Сохранено:", OUT, "| файлов:", len(all_files))
