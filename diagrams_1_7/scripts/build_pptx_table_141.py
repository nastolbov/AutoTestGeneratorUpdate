# -*- coding: utf-8 -*-
"""Презентация PowerPoint: Таблица 141 «Результаты системного тестирования»,
разбитая на 2 слайда — функциональные требования (1.7.2.1) и требования к
надёжности (1.7.2.2). Текст в ячейках сжат до презентационного вида (полные
формулировки остаются в docx-таблице). Выход: Системное_тестирование_таблица.pptx."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "Системное_тестирование_таблица.pptx")
FONT = "Times New Roman"
GRAY = RGBColor(0xD9, 0xD9, 0xD9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FUNC = [  # дата, кто, описание, ожидаемый результат
 ("19.05.2026", "Тестировщик", "Разбор XML-метамодели E3Core (StAX, NS_E/NS_E3/NS_MD)", "Метамодель разобрана в объектную модель (AppModel)"),
 ("19.05.2026", "Тестировщик", "Транслитерация русских имён в Java-идентификаторы", "Получены валидные уникальные идентификаторы"),
 ("19.05.2026", "Тестировщик", "Классификация сущностей (PRIMARY / CHILD / REFERENCE_DICTIONARY)", "Каждая сущность отнесена к верной категории"),
 ("19.05.2026", "Тестировщик", "Генерация Maven-проекта (pom.xml, общие классы, Page Object, тесты JUnit 5)", "Создан компилируемый проект автотестов"),
 ("20.05.2026", "Тестировщик", "Генерация полного набора UI-проверок (поля, валидация, CRUD, поиск, гриды)", "Сформированы тест-методы всех видов проверок"),
 ("20.05.2026", "Тестировщик", "Запуск автотестов через mvn test", "Прогон выполняется; формируются отчёты Surefire"),
 ("20.05.2026", "Тестировщик", "Сбор результатов из отчётов Surefire (StAX)", "Получена сводка «всего / успешно / провалено»"),
 ("20.05.2026", "Тестировщик", "Сохранение истории прогонов в SQLite", "Прогон и результаты сохранены в базе данных"),
 ("20.05.2026", "Разработчик", "Отображение истории прогонов в интерфейсе", "В журнале выведены все ранее сохранённые прогоны"),
 ("21.05.2026", "Разработчик", "HTML-отчёт с фотолетописью + CSV-выгрузка", "Сформированы HTML со скриншотами шагов и CSV"),
]
RELI = [
 ("21.05.2026", "Тестировщик", "Защита от некорректных действий и исходных данных", "Пустые/неверные параметры — сообщение, операция не запускается"),
 ("21.05.2026", "Тестировщик", "Обработка ошибки разбора XML (ParserException)", "Понятное сообщение, программа не завершается аварийно"),
 ("21.05.2026", "Разработчик", "Сохранение в БД одной транзакцией с откатом при ошибке", "При ошибке выполняется откат; частичные данные не сохраняются"),
 ("22.05.2026", "Разработчик", "Graceful degradation при сбое записи в БД", "Работа продолжается без сохранения истории; пользователь уведомлён"),
 ("22.05.2026", "Тестировщик", "Отсутствие вспомогательных средств (Apache Maven, Google Chrome)", "При прогоне выводится понятное сообщение об ошибке"),
 ("22.05.2026", "Разработчик", "Освобождение ресурсов (try-with-resources)", "Файловые потоки и JDBC-соединения корректно закрываются"),
 ("22.05.2026", "Тестировщик", "Сохранность БД между запусками + автосоздание", "БД создаётся при первом запуске; история не теряется"),
]

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def grid_style(tbl):
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    el = tblPr.find(qn("a:tableStyleId"))
    if el is None:
        el = tblPr.makeelement(qn("a:tableStyleId"), {}); tblPr.append(el)
    el.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # No Style, Table Grid


def set_cell(cell, text, size=10, bold=False, align=PP_ALIGN.LEFT, fill=WHITE):
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.05); cell.margin_right = Inches(0.05)
    cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.02)
    cell.fill.solid(); cell.fill.fore_color.rgb = fill
    tf = cell.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT
    r.font.color.rgb = RGBColor(0, 0, 0)


def add_slide(title, rows):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    tb = slide.shapes.add_textbox(Inches(0.4), Inches(0.18), Inches(12.5), Inches(0.7))
    pr = tb.text_frame.paragraphs[0]; run = pr.add_run(); run.text = title
    run.font.size = Pt(20); run.font.bold = True; run.font.name = FONT
    n = len(rows) + 1
    gf = slide.shapes.add_table(n, 5, Inches(0.35), Inches(1.0), Inches(12.63), Inches(0.42 * n))
    tbl = gf.table
    grid_style(tbl)
    for i, w in enumerate([1.05, 1.5, 4.78, 4.3, 1.0]):
        tbl.columns[i].width = Inches(w)
    headers = ["Дата", "Тестирование проводил", "Описание теста", "Ожидаемый результат", "Результат"]
    for j, h in enumerate(headers):
        set_cell(tbl.cell(0, j), h, size=11, bold=True, align=PP_ALIGN.CENTER, fill=GRAY)
    for i, (date, who, desc, exp) in enumerate(rows, start=1):
        set_cell(tbl.cell(i, 0), date, align=PP_ALIGN.CENTER)
        set_cell(tbl.cell(i, 1), who, align=PP_ALIGN.CENTER)
        set_cell(tbl.cell(i, 2), desc)
        set_cell(tbl.cell(i, 3), exp)
        set_cell(tbl.cell(i, 4), "Успех", align=PP_ALIGN.CENTER)
    for row in tbl.rows:
        row.height = Inches(0.3)


add_slide("Результаты системного тестирования: функциональные требования (п. 1.7.2.1)", FUNC)
add_slide("Результаты системного тестирования: требования к надёжности (п. 1.7.2.2)", RELI)
prs.save(OUT)
print("Сохранено:", os.path.abspath(OUT))
print(f"слайдов: {len(prs.slides._sldIdLst)} | функц.: {len(FUNC)} | надёжность: {len(RELI)}")
