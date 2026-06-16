# -*- coding: utf-8 -*-
"""Переписать §3.3 «Условия эксплуатации» в ТЗ (Приложение 1) Столбов_ВКР.docx
по образцу другого диплома: перечень необходимых средств + базовые знания
пользователя + перечень экранных форм интерфейса (адаптировано под наш
Java/JavaFX/E3Core-проект вместо 1С)."""
import os
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

DOC = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "Столбов_ВКР.docx")
FONT = "Times New Roman"
d = docx.Document(DOC)
ps = d.paragraphs


def idx(starts):
    for i, p in enumerate(ps):
        if p.text.strip().startswith(starts):
            return i
    raise SystemExit("не найдено: " + starts)


i33 = idx("3.3 Условия эксплуатации")
i34 = idx("3.4 Требование к составу")
# якорь — пустой абзац перед 3.4 (или сам 3.4)
anchor = ps[i34 - 1] if ps[i34 - 1].text.strip() == "" else ps[i34]
base_style = ps[i33].style

print("Было §3.3:")
for i in range(i33, i34):
    if ps[i].text.strip():
        print("  ", ps[i].text.strip())

# удалить старое тело §3.3 (между заголовком и якорем)
for p in ps[i33 + 1:i34]:
    if p is anchor:
        continue
    p._element.getparent().remove(p._element)


def _f(run):
    run.font.name = FONT; run.font.size = Pt(14)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def add_plain(text):
    np = anchor.insert_paragraph_before(); np.style = base_style
    np.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    np.paragraph_format.first_line_indent = Cm(1.25)
    np.paragraph_format.line_spacing = 1.5
    _f(np.add_run(text))


def add_bullet(text):
    np = anchor.insert_paragraph_before(); np.style = base_style
    np.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    np.paragraph_format.left_indent = Cm(1.25)
    np.paragraph_format.first_line_indent = Cm(-0.5)
    np.paragraph_format.line_spacing = 1.5
    _f(np.add_run("•")); _f(np.add_run("\t" + text))


add_plain("Для работы программы потребуется:")
for t in ["комплект разработки JDK Java SE 17 или новее;",
          "система сборки Apache Maven версии 3.8 и выше;",
          "браузер Google Chrome актуальной версии."]:
    add_bullet(t)
add_plain("Пользователь должен обладать базовыми знаниями операционной системы и принципов "
          "автоматизированного тестирования веб-приложений на платформе E3Core. Ему будет "
          "предложен графический интерфейс программы, который содержит следующие экранные формы:")
for t in ["главное окно с панелью параметров (XML-файл метамодели, URL сайта, логин, пароль, "
          "каталог вывода, режим headless);",
          "дерево сущностей метамодели с результатами классификации (PRIMARY, CHILD, "
          "REFERENCE_DICTIONARY);",
          "таблица результатов прогона тестов (класс, метод, статус, длительность);",
          "форма выбора сущностей и видов проверок для запуска;",
          "журнал (история) прогонов."]:
    add_bullet(t)

d.save(DOC)

d2 = docx.Document(DOC)
p2 = d2.paragraphs
j = next(i for i, p in enumerate(p2) if p.text.strip().startswith("3.3 Условия"))
k = next(i for i, p in enumerate(p2) if p.text.strip().startswith("3.4 Требование к составу"))
print("\nСтало §3.3:")
for i in range(j, k):
    if p2[i].text.strip():
        print("  ", p2[i].text.strip())
print("\nВсего абзацев:", len(p2), "| таблиц:", len(d2.tables))
