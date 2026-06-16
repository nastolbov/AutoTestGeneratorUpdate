# -*- coding: utf-8 -*-
"""Исправление недочётов ТЗ (Приложение 1) прямо в основном дипломе
Столбов_ВКР.docx: технические характеристики (п. 3.4) приводятся к значениям
п. 1.7.2.4 (≥2 ГГц / 4 ГБ (рек. 8) / 2 ГБ / 1024×768), а требование об уровнях
тестов переформулируется на «по умолчанию FULL». Правка точечная, форматирование
сохраняется; перед каждой заменой проверяется текущий текст абзаца (assert)."""
import os
from docx import Document

DOC = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "Столбов_ВКР.docx")
d = Document(DOC)
ps = d.paragraphs

n0_par, n0_tbl = len(ps), len(d.tables)


def rewrite_keep_first(p, new_text):
    """Весь текст абзаца — в первый run (его формат сохраняется), остальные очищаются."""
    assert p.runs, "пустой абзац"
    p.runs[0].text = new_text
    for r in p.runs[1:]:
        r.text = ""


def rewrite_bullet(p, new_text):
    """Маркированный пункт: run[0]='•' сохраняется, текст — во второй run."""
    p.runs[1].text = "\t" + new_text
    for r in p.runs[2:]:
        r.text = ""


def expect(i, needle):
    got = ps[i].text.strip()
    assert needle in got, f"P{i}: ожидали «{needle}», получили «{got[:90]}»"


# --- п. 3.4 «состав и параметры технических средств» ---
expect(1541, "Рекомендуемая конфигурация")
rewrite_keep_first(ps[1541], "Минимальная конфигурация:")

expect(1542, "не менее 2,3 ГГц")
rewrite_keep_first(ps[1542], "Процессор с тактовой частотой не менее 2 ГГц.")

expect(1543, "ОЗУ не менее 8")
rewrite_keep_first(ps[1543], "Объём ОЗУ не менее 4 ГБ (рекомендуется 8 ГБ).")

expect(1544, "жёсткого диска не менее 50")
rewrite_keep_first(ps[1544], "Объём жёсткого диска не менее 2 ГБ.")

expect(1545, "1920")
rewrite_keep_first(ps[1545], "Монитор, поддерживающий разрешение не менее 1024×768 точек.")

# --- список функций: уровни тестов ---
expect(1514, "поддержка трёх уровней тестов")
rewrite_bullet(ps[1514],
    "генерация набора проверок по уровням: SMOKE (наличие полей), "
    "BASIC (дополнительно поиск и гриды), FULL (дополнительно полный CRUD-цикл); "
    "по умолчанию применяется уровень FULL;")

d.save(DOC)

# --- проверка целостности ---
d2 = Document(DOC)
p2 = d2.paragraphs
assert len(p2) == n0_par, f"число абзацев изменилось: {len(p2)} != {n0_par}"
assert len(d2.tables) == n0_tbl, f"число таблиц изменилось: {len(d2.tables)} != {n0_tbl}"
print("OK. абзацев:", len(p2), "таблиц:", len(d2.tables))
for i in (1514, 1541, 1542, 1543, 1544, 1545):
    print(f"  P{i}: {p2[i].text.strip()}")
