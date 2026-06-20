# -*- coding: utf-8 -*-
"""Таблица 141 «Результаты системного тестирования на соответствие требованиям ТЗ».
Покрывает КАЖДОЕ требование из п. 1.7.2.1 (10 функциональных) и п. 1.7.2.2
(7 надёжности) — по одной строке на требование, в порядке требований (трассируемость).
Формат — 5 столбцов, как в дипломе. Выход: Таблица_141_системное.docx."""
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "Таблица_141_системное.docx")
FONT = "Times New Roman"
doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT; st.font.size = Pt(14)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)
sec = doc.sections[0]
sec.left_margin = Cm(3); sec.right_margin = Cm(1.5); sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)


def _f(run, size=14, bold=False):
    run.font.name = FONT; run.font.size = Pt(size); run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def _shade(cell, hexc="D9D9D9"):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexc); tcPr.append(sh)


def _ct(cell, text, bold=False, size=12, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""; p = cell.paragraphs[0]; p.alignment = align
    p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
    _f(p.add_run(text), size, bold=bold)


# подпись
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
_f(p.add_run("Таблица 141"), 14)
p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER; p2.paragraph_format.space_after = Pt(2)
_f(p2.add_run("Результаты системного тестирования на соответствие требованиям ТЗ"), 14)

W = "Тестировщик"
D = "Разработчик"
# (дата, кто, описание теста = требование, ожидаемый результат)
ROWS = [
 # ---- функциональные требования (п. 1.7.2.1) ----
 ("19.05.2026", W, "Разбор XML-метамодели формата E3Core потоковым парсером StAX с поддержкой пространств имён NS_E, NS_E3 и NS_MD",
  "Эталонная метамодель подсистемы E3Core полностью разобрана в объектную модель (AppModel): сущности, свойства, операции и поиски извлечены корректно"),
 ("19.05.2026", W, "Транслитерация русских имён сущностей и атрибутов в идентификаторы языка Java",
  "Имена преобразованы в корректные (валидные) Java-идентификаторы без потери уникальности"),
 ("19.05.2026", W, "Классификация сущностей модели по категориям PRIMARY, CHILD и REFERENCE_DICTIONARY",
  "Каждая сущность отнесена к верной категории; результат отражён в дереве сущностей и CSV-отчёте о классификации"),
 ("19.05.2026", W, "Генерация Maven-проекта автотестов: pom.xml, общие классы, Page Object и тест-класс JUnit 5 для каждой PRIMARY-сущности",
  "В каталоге вывода создан компилируемый Maven-проект с pom.xml, общими классами и Page Object / тест-классом для каждой PRIMARY-сущности"),
 ("20.05.2026", W, "Генерация полного набора UI-проверок: наличие полей формы, валидация обязательных полей, создание/изменение/удаление (CRUD), логическое удаление и архивирование, параметрический поиск и проверка гридов",
  "Для тестируемых сущностей сформированы тест-методы всех предусмотренных видов проверок в соответствии с операциями метамодели"),
 ("20.05.2026", W, "Запуск сгенерированных автотестов через команду mvn test",
  "Прогон автотестов запускается и выполняется до завершения; формируются XML-отчёты Surefire"),
 ("20.05.2026", W, "Сбор результатов прогона из XML-отчётов Surefire потоковым парсером StAX",
  "Сводка «всего / успешно / провалено» и сведения по каждому тесту получены из отчётов корректно"),
 ("20.05.2026", W, "Сохранение истории прогонов в локальную базу данных SQLite",
  "Сведения о прогоне и результаты тестов сохранены в базе данных и доступны для последующего чтения"),
 ("20.05.2026", D, "Отображение истории прогонов в графическом интерфейсе",
  "В журнале прогонов выведены все ранее сохранённые прогоны с их показателями"),
 ("21.05.2026", D, "Формирование пользовательского HTML-отчёта с фотолетописью (скриншоты на каждом шаге теста) и CSV-выгрузки",
  "Сформированы HTML-отчёт со скриншотами шагов и CSV-файл с результатами прогона"),
 # ---- требования к надёжности (п. 1.7.2.2) ----
 ("21.05.2026", W, "Защита от некорректных действий оператора и ошибочных исходных данных (пустые или неверные параметры)",
  "При отсутствии обязательных параметров выводится понятное сообщение, операция не запускается"),
 ("21.05.2026", W, "Обработка ошибки разбора XML: формирование исключения ParserException",
  "При некорректном XML формируется ParserException с понятным сообщением, программа не завершается аварийно"),
 ("21.05.2026", D, "Сохранение результатов прогона в локальную базу данных в рамках одной транзакции с откатом при ошибке",
  "Запись выполняется как единая транзакция; при ошибке выполняется откат, частичные данные не сохраняются, целостность БД не нарушается"),
 ("22.05.2026", D, "Поведение при сбое записи в базу данных (graceful degradation)",
  "Программа продолжает работу без сохранения истории прогона и уведомляет пользователя о сбое"),
 ("22.05.2026", W, "Поведение при отсутствии вспомогательных средств (Apache Maven, Google Chrome)",
  "При попытке прогона тестов выводится понятное сообщение об ошибке; программа не завершается аварийно"),
 ("22.05.2026", D, "Освобождение файловых потоков и JDBC-соединений через конструкцию try-with-resources",
  "После операций все открытые ресурсы (файлы, соединения с БД) корректно закрываются"),
 ("22.05.2026", W, "Сохранность локальной базы данных истории между запусками и её автоматическое создание при первом запуске",
  "При первом запуске база данных и схема создаются автоматически; история прогонов сохраняется между запусками программы"),
]

headers = ["Дата", "Тестирование\nпроводил", "Описание теста", "Ожидаемый результат", "Результат\nтестирования"]
widths = [2.0, 2.3, 5.5, 4.3, 1.9]
t = doc.add_table(rows=2 + len(ROWS), cols=5)
t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
for j, h in enumerate(headers):
    _ct(t.rows[0].cells[j], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER); _shade(t.rows[0].cells[j])
for j in range(5):
    _ct(t.rows[1].cells[j], str(j + 1), align=WD_ALIGN_PARAGRAPH.CENTER)
for i, (date, who, desc, exp) in enumerate(ROWS, start=2):
    _ct(t.rows[i].cells[0], date, align=WD_ALIGN_PARAGRAPH.CENTER)
    _ct(t.rows[i].cells[1], who)
    _ct(t.rows[i].cells[2], desc)
    _ct(t.rows[i].cells[3], exp)
    _ct(t.rows[i].cells[4], "Успех", align=WD_ALIGN_PARAGRAPH.CENTER)
for i in range(len(ROWS) + 2):
    for j, wd in enumerate(widths):
        t.rows[i].cells[j].width = Cm(wd)

doc.save(OUT)
nf = sum(1 for r in ROWS[:10]); nr = len(ROWS) - 10
print("Сохранено:", os.path.abspath(OUT))
print(f"строк требований: {len(ROWS)} (функциональных {nf} + надёжности {nr})")
