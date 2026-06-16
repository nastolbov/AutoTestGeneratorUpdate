# -*- coding: utf-8 -*-
"""Выравнивание раздела 1.7 в основном дипломе Столбов_ВКР.docx по факту кода:
строка системной таблицы, проверявшая несуществующее требование «три уровня
тестов SMOKE/BASIC/FULL», и упоминание «уровень FULL» в таблице автономного
тестирования заменяются на «полный набор UI-проверок». Перед правкой
проверяется текущий текст ячейки (assert)."""
import os
from docx import Document

DOC = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "Столбов_ВКР.docx")
d = Document(DOC)
n0_par, n0_tbl = len(d.paragraphs), len(d.tables)

FULL_SET = "полного набора UI-проверок (поля, валидация, CRUD, поиск, гриды)"


def set_cell_para(cell, new_text, must_contain):
    p = cell.paragraphs[0]
    assert must_contain in p.text, f"ожидали «{must_contain}», получили «{p.text[:80]}»"
    p.runs[0].text = new_text
    for r in p.runs[1:]:
        r.text = ""


# T193 r7c1 — системное тестирование требований ТЗ
set_cell_para(d.tables[193].rows[7].cells[1],
              "Генерация для каждой сущности " + FULL_SET,
              "Поддержка трёх уровней тестов")

# T191 r14c4 — автономное тестирование (TestClassWriter)
set_cell_para(d.tables[191].rows[14].cells[4],
              "Успех, для сущности «Гаражно-строительный кооператив» сгенерирован полный набор "
              "тестов (поля, валидация, CRUD, поиск, гриды)",
              "уровнем")

d.save(DOC)

# проверка
d2 = Document(DOC)
assert len(d2.paragraphs) == n0_par and len(d2.tables) == n0_tbl, "целостность нарушена"
print("OK. абзацев:", len(d2.paragraphs), "таблиц:", len(d2.tables))
print("T193 r7c1:", d2.tables[193].rows[7].cells[1].text.strip())
print("T191 r14c4:", d2.tables[191].rows[14].cells[4].text.strip())
